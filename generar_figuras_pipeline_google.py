#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Regenera, con el mismo estilo matplotlib del resto del pipeline de la tesis,
las figuras que hasta ahora eran exportaciones antiguas de Google Sheets:

  - "Figura 5.1 - SQL.png"                 (radar Component Matching)
  - "Figura 5.3 - con autocorreccion.png"  (radar Execution / Sin errores / Keywords)
  - "Tabla 4 Time To Answer.png"           (tabla de TTA por configuracion)
  - "Figura 5.9 Correlacion sin AC.png"    (dispersion SIN_AUTO vs TTA promedio)
  - "Figura - Evaluacion de CRR.png"       (CRR a escala del dataset multiturno)

Todos los numeros de un solo turno se calculan sobre el MISMO subconjunto de
53 consultas gold (para la config completa se excluyen los 759 turnos
multiturno, que comparten Parametros pero tienen otro texto de pregunta).
El CRR se calcula sobre los 313 turnos de seguimiento dependientes de memoria
de la corrida real multiturno.
"""
import os, csv, json, statistics as st
from collections import defaultdict

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter
from matplotlib.patches import Patch

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "dataset_sql_40_med_complex.xlsx")
AP = os.path.join(BASE, "analisis_pruebas")

# ----- paleta / estilo (coherente con generar_figuras_ranking_tesis.py) -----
C_BLUE, C_GREEN, C_ORANGE = "#2a78d6", "#008300", "#eb6834"
C_TEAL, C_YELLOW, C_MAGENTA = "#1f6f78", "#eda100", "#e87ba4"
INK, INK2, SURF = "#0b0b0b", "#52514e", "#fcfcfb"
GRID = "#e4e2dd"
plt.rcParams.update({
    "figure.facecolor": SURF, "axes.facecolor": SURF, "savefig.facecolor": SURF,
    "font.family": "DejaVu Sans", "font.size": 11,
    "axes.edgecolor": "#c9c7c2", "axes.labelcolor": INK,
    "text.color": INK, "xtick.color": INK2, "ytick.color": INK2,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8,
    "axes.spines.top": False, "axes.spines.right": False,
    "figure.dpi": 150, "savefig.dpi": 200, "savefig.bbox": "tight",
})


# --------------------------------------------------------------------------- #
# 1. Datos de un solo turno (53 gold) directamente del Excel
# --------------------------------------------------------------------------- #
def num(v):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def load_ablacion():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    ws = wb["Medianas_y_Complejas"]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    H = {h: i for i, h in enumerate(hdr)}
    allrows = [dict(zip(hdr, r)) for r in it
               if r[H["chatInput"]] and str(r[H["EVALUADO"]]) == "1"]
    wb.close()
    gold53 = {str(r["chatInput"]).strip() for r in allrows
              if r["Parametros"] == "[Rag],[Din],[Dea]"}
    seen, data = set(), {}
    for r in allrows:
        q = str(r["chatInput"]).strip()
        if q not in gold53:
            continue
        p = r["Parametros"]
        if (p, q) in seen:
            continue
        seen.add((p, q))
        d = data.setdefault(p, {k: [] for k in
                                ("sel", "whe", "grp", "kw", "ex", "err", "sa", "tta")})
        for k, col in [("sel", "VAL_SELECT"), ("whe", "VAL_WHERE"),
                       ("grp", "VAL_GROUP"), ("kw", "VAL_KEYWORDS"),
                       ("ex", "VAL_EXECUTION"), ("err", "VAL_ERRORES"),
                       ("sa", "SIN_AUTO_CORRECCION")]:
            v = num(r.get(col))
            if v is not None:
                d[k].append(v)
        t = num(r.get("N8nTimeToResponse"))
        if t is not None:
            d["tta"].append(t / 1000.0)
    out = {}
    m = lambda xs: st.mean(xs) * 100 if xs else None
    for p, d in data.items():
        out[p] = {
            "n": len(d["ex"]),
            "SELECT": m(d["sel"]), "WHERE": m(d["whe"]),
            "GROUP": m(d["grp"]), "KEYWORDS": m(d["kw"]),
            "EX": m(d["ex"]), "ERR": m(d["err"]), "SA": m(d["sa"]),
            "tta_prom": st.mean(d["tta"]) if d["tta"] else None,
            "tta_min": min(d["tta"]) if d["tta"] else None,
            "tta_max": max(d["tta"]) if d["tta"] else None,
        }
    return out


# --------------------------------------------------------------------------- #
# Radar helper
# --------------------------------------------------------------------------- #
def radar(ax, axes_labels, series, colors):
    """series: list of (label, [values 0..100])."""
    N = len(axes_labels)
    ang = np.linspace(0, 2 * np.pi, N, endpoint=False).tolist()
    ang += ang[:1]
    ax.set_theta_offset(np.pi / 2)
    ax.set_theta_direction(-1)
    ax.set_xticks(ang[:-1])
    ax.set_xticklabels(axes_labels, fontsize=10, color=INK)
    ax.set_ylim(0, 100)
    ax.set_yticks([25, 50, 75, 100])
    ax.set_yticklabels(["25", "50", "75", "100"], fontsize=8, color=INK2)
    ax.set_rlabel_position(90)
    ax.grid(color=GRID, linewidth=0.9)
    ax.spines["polar"].set_edgecolor("#c9c7c2")
    for (lab, vals), c in zip(series, colors):
        v = list(vals) + [vals[0]]
        ax.plot(ang, v, color=c, linewidth=2.1, label=lab, zorder=3)
        ax.fill(ang, v, color=c, alpha=0.06, zorder=2)


def bracket_label(p):
    return p  # se conservan las etiquetas [Rag],[Memoria],... como en la tesis


def add_table(fig, rect, col_labels, cell_rows, col_widths=None):
    ax = fig.add_axes(rect)
    ax.axis("off")
    tbl = ax.table(cellText=cell_rows, colLabels=col_labels,
                   cellLoc="center", loc="center", colWidths=col_widths)
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(8.6)
    tbl.scale(1, 1.35)
    for (r, c), cell in tbl.get_celld().items():
        cell.set_edgecolor("#d7d5cf")
        cell.set_linewidth(0.7)
        if r == 0:
            cell.set_facecolor("#eef2f6")
            cell.set_text_props(weight="bold", color=INK)
        else:
            cell.set_facecolor(SURF if r % 2 else "#f7f6f3")
            if c == 0:
                cell.set_text_props(color=INK)
    return tbl


# --------------------------------------------------------------------------- #
# 2. Figura 5.1 - radar Component Matching
# --------------------------------------------------------------------------- #
def fig_51(D):
    axes = ["SELECT", "WHERE", "GROUP BY", "KEYWORDS"]
    keys = ["SELECT", "WHERE", "GROUP", "KEYWORDS"]
    configs = ["[Rag],[Memoria],[Intension],[Din],[Dea]",
               "[Intension],[Din],[Dea]",
               "[Rag],[Din],[Dea]",
               "[Rag]"]
    colors = [C_BLUE, C_TEAL, C_ORANGE, C_GREEN]
    fig = plt.figure(figsize=(8.4, 7.5))
    ax = fig.add_axes([0.12, 0.37, 0.76, 0.50], polar=True)
    series = [(bracket_label(c), [D[c][k] for k in keys]) for c in configs]
    radar(ax, axes, series, colors)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, 1.24), ncol=2,
              frameon=False, fontsize=9.2)
    rows = [[bracket_label(c)] + [f"{D[c][k]:.2f}\\%".replace("\\", "")
                                  for k in keys] for c in configs]
    add_table(fig, [0.02, 0.03, 0.96, 0.26],
              ["Parámetros", "SELECT", "WHERE", "GROUP BY", "KEYWORDS"], rows,
              col_widths=[0.42, 0.145, 0.145, 0.145, 0.145])
    fig.savefig(os.path.join(BASE, "Figura 5.1 - SQL.png"))
    plt.close(fig)
    print("  -> Figura 5.1 - SQL.png")


# --------------------------------------------------------------------------- #
# 3. Figura 5.3 - radar Execution / Sin errores / Keywords
# --------------------------------------------------------------------------- #
def fig_53(D):
    axes = ["EXECUTION", "SIN ERRORES", "KEYWORDS"]
    keys = ["EX", "ERR", "KEYWORDS"]
    configs = ["[Intension],[Din],[Dea]",
               "[Rag],[Din],[Dea]",
               "[Rag],[Dea]",
               "[Rag],[Memoria]"]
    colors = [C_TEAL, C_BLUE, C_ORANGE, C_GREEN]
    fig = plt.figure(figsize=(8.4, 7.5))
    ax = fig.add_axes([0.14, 0.37, 0.72, 0.50], polar=True)
    series = [(bracket_label(c), [D[c][k] for k in keys]) for c in configs]
    radar(ax, axes, series, colors)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, 1.24), ncol=2,
              frameon=False, fontsize=9.2)
    rows = [[bracket_label(c)] + [f"{D[c][k]:.2f}\\%".replace("\\", "")
                                  for k in keys] for c in configs]
    add_table(fig, [0.04, 0.03, 0.92, 0.25],
              ["Parámetros", "EXECUTION", "SIN ERRORES", "KEYWORDS"], rows,
              col_widths=[0.40, 0.20, 0.20, 0.20])
    fig.savefig(os.path.join(BASE, "Figura 5.3 - con autocorreccion.png"))
    plt.close(fig)
    print("  -> Figura 5.3 - con autocorreccion.png")


# --------------------------------------------------------------------------- #
# 4. Tabla 4 - Time To Answer
# --------------------------------------------------------------------------- #
NICE = {"Rag": "RAG", "Memoria": "Memoria", "Intension": "Intensión",
        "Din": "DIN", "Dea": "DEA"}


def nice(p):
    return " + ".join(NICE[t.strip("[]")] for t in p.split(","))


def fig_tabla4(D):
    order = sorted(D.keys(), key=lambda p: -(D[p]["tta_prom"] or 0))
    col = ["Configuración (módulos activos)", "TTA prom. (s)", "TTA mín. (s)",
           "TTA máx. (s)", "Execution (%)", "Sin errores (%)",
           "Sin autocorrección (%)"]
    rows = []
    for p in order:
        d = D[p]
        rows.append([nice(p), f"{d['tta_prom']:.1f}", f"{d['tta_min']:.1f}",
                     f"{d['tta_max']:.1f}", f"{d['EX']:.2f}",
                     f"{d['ERR']:.2f}", f"{d['SA']:.2f}"])
    fig = plt.figure(figsize=(9.6, 0.42 * len(rows) + 1.0))
    ax = fig.add_axes([0.01, 0.01, 0.98, 0.98])
    ax.axis("off")
    tbl = ax.table(cellText=rows, colLabels=col, cellLoc="center", loc="center")
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(8.4)
    tbl.scale(1, 1.32)
    full = "[Rag],[Memoria],[Intension],[Din],[Dea]"
    best_tta = order[-1]
    for (r, c), cell in tbl.get_celld().items():
        cell.set_edgecolor("#d7d5cf")
        cell.set_linewidth(0.6)
        if r == 0:
            cell.set_facecolor("#eef2f6")
            cell.set_text_props(weight="bold", color=INK)
            cell.set_height(0.09)
        else:
            p = order[r - 1]
            base = SURF if r % 2 else "#f7f6f3"
            if p == full:
                base = "#fdece2"       # config completa (conversacional)
            cell.set_facecolor(base)
            if c == 0:
                cell.set_text_props(ha="left", color=INK)
                cell.PAD = 0.03
    tbl.auto_set_column_width(col=list(range(len(col))))
    fig.savefig(os.path.join(BASE, "Tabla 4 Time To Answer.png"))
    plt.close(fig)
    print("  -> Tabla 4 Time To Answer.png")


# --------------------------------------------------------------------------- #
# 5. Figura 5.9 - dispersion SIN_AUTO vs TTA promedio
# --------------------------------------------------------------------------- #
def fig_59(D):
    xs = [D[p]["tta_prom"] for p in D]
    ys = [D[p]["SA"] for p in D]
    full = "[Rag],[Memoria],[Intension],[Din],[Dea]"
    fig, ax = plt.subplots(figsize=(8.4, 5.8))
    # tendencia
    z = np.polyfit(xs, ys, 1)
    xline = np.linspace(min(xs) - 1, max(xs) + 1, 50)
    ax.plot(xline, np.polyval(z, xline), color="#b9b7b1", linewidth=1.6,
            linestyle="--", zorder=1)
    r = np.corrcoef(xs, ys)[0, 1]
    for p in D:
        c = C_ORANGE if p == full else C_BLUE
        ax.scatter(D[p]["tta_prom"], D[p]["SA"], s=78, color=c,
                   edgecolor="white", linewidth=1.1, zorder=3)
    ann = {full: (6, 6), "[Rag],[Din],[Dea]": (6, -14),
           "[Intension],[Din],[Dea]": (-4, 8), "[Rag],[Dea]": (6, 4)}
    for p, off in ann.items():
        if p in D:
            ax.annotate(nice(p), (D[p]["tta_prom"], D[p]["SA"]),
                        textcoords="offset points", xytext=off, fontsize=8.4,
                        color=INK)
    ax.set_xlabel("Time-To-Answer promedio (s)")
    ax.set_ylabel("SIN_AUTO_CORRECCION")
    ax.yaxis.set_major_formatter(PercentFormatter())
    ax.set_title(f"Correlación entre generación al primer intento y TTA "
                 f"(r = {r:.2f})", loc="left", fontsize=12.5, pad=12)
    ax.legend(handles=[Patch(color=C_ORANGE,
                             label="Configuración completa (conversacional)")],
              frameon=False, loc="upper right", fontsize=8.8)
    fig.savefig(os.path.join(BASE, "Figura 5.9 Correlacion sin AC.png"))
    plt.close(fig)
    print("  -> Figura 5.9 Correlacion sin AC.png")


# --------------------------------------------------------------------------- #
# 6. Figura CRR a escala del dataset multiturno
# --------------------------------------------------------------------------- #
def load_crr():
    design = list(csv.DictReader(
        open(os.path.join(AP, "preguntas_multiturno.csv"), encoding="utf-8-sig")))
    run = {int(r["excel_id"]): r for r in json.load(
        open(os.path.join(AP, "multiturn_run", "resumen_metricas.json"),
             encoding="utf-8"))}
    LBL = {"ctx_top5": "Top-N (“los primeros N”)",
           "ctx_total": "Total / conteo", "ctx_max": "Máximo / mínimo",
           "ctx_orden": "Reordenamiento", "ctx_promedio": "Promedio"}
    agg = defaultdict(lambda: [0, 0])
    for d in design:
        t = d["turn_type"]
        if not t.startswith("ctx_"):
            continue
        rid = int(d["excel_id"])
        if rid not in run:
            continue
        agg[t][1] += 1
        if run[rid]["action"] == d["expected_action"]:
            agg[t][0] += 1
    rows = [(LBL.get(t, t), h, n) for t, (h, n) in agg.items()]
    rows.sort(key=lambda r: -r[2])
    tot_h = sum(r[1] for r in rows)
    tot_n = sum(r[2] for r in rows)
    return rows, tot_h, tot_n


def fig_crr():
    rows, tot_h, tot_n = load_crr()
    labels = [r[0] for r in rows]
    vals = [r[1] / r[2] * 100 for r in rows]
    ns = [(r[1], r[2]) for r in rows]
    y = np.arange(len(labels))[::-1]
    fig, ax = plt.subplots(figsize=(8.8, 4.2))
    bars = ax.barh(y, vals, color=C_BLUE, height=0.62, zorder=3)
    for yi, v, (h, n) in zip(y, vals, ns):
        ax.text(min(v + 1.2, 88), yi, f"{v:.1f}%  ({h}/{n})",
                va="center", fontsize=9.5,
                color="white" if v > 92 else INK2,
                ha="right" if v > 92 else "left")
        if v > 92:
            ax.text(v - 1.2, yi, f"{v:.1f}%  ({h}/{n})", va="center",
                    ha="right", fontsize=9.5, color="white", fontweight="bold")
    # limpiar dobles etiquetas: volver a pintar limpio
    ax.clear()
    bars = ax.barh(y, vals, color=C_BLUE, height=0.62, zorder=3)
    for yi, v, (h, n) in zip(y, vals, ns):
        inside = v > 55
        ax.text(v - 1.5 if inside else v + 1.5, yi, f"{v:.1f}%  ({h}/{n})",
                va="center", ha="right" if inside else "left", fontsize=9.6,
                color="white" if inside else INK2,
                fontweight="bold" if inside else "normal")
    ax.set_yticks(y)
    ax.set_yticklabels(labels, fontsize=10)
    ax.set_xlim(0, 112)
    ax.set_xticks([0, 25, 50, 75, 100])
    ax.xaxis.set_major_formatter(PercentFormatter())
    ax.grid(axis="y", visible=False)
    ax.set_xlabel("Context Retention Rate (turnos con reutilización correcta de contexto)")
    ax.set_title(f"CRR por tipo de seguimiento contextual — {tot_h}/{tot_n} "
                 f"= {tot_h/tot_n*100:.1f}% global",
                 loc="left", fontsize=12.5, pad=12)
    fig.savefig(os.path.join(BASE, "Figura - Evaluacion de CRR.png"))
    plt.close(fig)
    print(f"  -> Figura - Evaluacion de CRR.png ({tot_h}/{tot_n} = "
          f"{tot_h/tot_n*100:.1f}%)")


def main():
    print("Regenerando figuras (datos actuales, estilo matplotlib):")
    D = load_ablacion()
    fig_51(D)
    fig_53(D)
    fig_tabla4(D)
    fig_59(D)
    fig_crr()
    print("Listo.")


if __name__ == "__main__":
    main()
